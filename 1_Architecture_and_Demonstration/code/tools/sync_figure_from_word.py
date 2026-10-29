"""Keep the vector figure text aligned with the researcher's edited Word file."""
from pathlib import Path
import argparse, json
from docx import Document
ROOT = Path(__file__).resolve().parents[2]

def sync(source):
    document = Document(source); table = document.tables[0]
    target = ROOT / 'figures/Figure_Content.json'
    spec = json.loads(target.read_text(encoding='utf-8'))
    main = {}
    for label, column in [('A1', 0), ('A2', 2)]:
        boxes = []
        for row in [1, 3, 5, 7, 9, 11]:
            paragraphs = [p.text.strip() for p in table.rows[row].cells[column].paragraphs if p.text.strip()]
            boxes.append([paragraphs[0], [s.lstrip('• ').strip() for s in paragraphs[1:]]])
        main[label] = dict(header=[table.rows[0].cells[column].text.strip()], boxes=boxes)
    assert [b[0] for b in main['A1']['boxes']][1::3] == ['Python script A1-1', 'Python script A1-2']
    assert [b[0] for b in main['A2']['boxes']][1::3] == ['Python script A2-1', 'Python script A2-2']
    spec['main'] = main
    spec['script_numbering'] = 'Each main numbered box maps to the runnable file with the same label: A1-1.py, A1-2.py, A2-1.py, A2-2.py. Tutorial labels map separately to the existing tutorial implementation.'
    target.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding='utf-8')
    return spec

if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--input', type=Path, default=ROOT/'figures/Architecture_Figure_Editable.docx')
    args = p.parse_args(); sync(args.input); print(str(ROOT/'figures/Figure_Content.json'))
