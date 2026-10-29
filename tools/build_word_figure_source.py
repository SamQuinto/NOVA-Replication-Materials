"""Native Word text and boxes for editing the current architecture comparison."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
W=Path(__file__).resolve().parents[1];OUT=W/'figures'
def borders(cell,visible=True):
    pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single' if visible else 'nil');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'000000');b.append(el)
    pr.append(b);m=OxmlElement('w:tcMar')
    for side,value in [('top',75),('left',100),('bottom',75),('right',100)]:
        el=OxmlElement('w:'+side);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');m.append(el)
    pr.append(m)
def main():
    spec=json.loads((OUT/'Figure_Content.json').read_text(encoding='utf-8'))['main'];doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.45);s.left_margin=s.right_margin=Inches(.45)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(10)
        st.paragraph_format.space_after=Pt(2);st.paragraph_format.line_spacing=1
    title=doc.add_paragraph('NOVA scorer architectures',style='Title');title.runs[0].font.size=Pt(14)
    # Remove the default template's blue Title rule; retain plain black text.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for paragraph in doc.paragraphs:
        for border in list(paragraph._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
    table=doc.add_table(rows=1,cols=3);table.autofit=False
    widths=[3.65,.3,3.65]
    for c,width in zip(table.columns,widths):c.width=Inches(width)
    for i,key in enumerate(['A1','A2']):
        cell=table.rows[0].cells[i*2];borders(cell,False);p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run('\n'.join(spec[key]['header']));r.bold=True;r.font.size=Pt(11)
    borders(table.rows[0].cells[1],False)
    for step in range(6):
        if step:
            row=table.add_row()
            for j,cell in enumerate(row.cells):
                borders(cell,False)
                if j!=1:
                    p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('↓');r.font.size=Pt(16)
        row=table.add_row()
        for i,key in enumerate(['A1','A2']):
            cell=row.cells[i*2];borders(cell);title,body=spec[key]['boxes'][step];p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(title);r.bold=True;r.font.size=Pt(10.5)
            for item in body:
                p=cell.add_paragraph();p.paragraph_format.space_after=Pt(1);p.add_run('• '+item)
            # Do not split a diagram box over pages.
            trpr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');trpr.append(keep)
        borders(row.cells[1],False)
    p=doc.add_paragraph('Script numbers identify processing stages, not separate source files. The optional second tutorial reading is shown in its separate diagram.');p.paragraph_format.space_before=Pt(5)
    for r in p.runs:r.font.size=Pt(9)
    doc.core_properties.title='NOVA scorer architectures';doc.core_properties.author='';doc.save(OUT/'Architecture_Figure_Editable.docx')
    print(str(OUT/'Architecture_Figure_Editable.docx'))
if __name__=='__main__':main()
